"""Excel builders. Excel can't embed fonts - deliver a PDF too if typography matters."""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from brand import BRAND


def _fill(h):
    return PatternFill("solid", fgColor=h.lstrip("#"))


def add_title_block(ws, title, subtitle=""):
    """Returns the row number the table header should start on."""
    ws["A1"] = title
    ws["A1"].font = Font(name=BRAND.fonts["heading"], size=16, color=BRAND.BLUE.lstrip("#"))
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=BRAND.fonts["body"], size=11, color=BRAND.INK.lstrip("#"))
    return 4


def style_sheet(ws, header_row, number_format="#,##0"):
    for cell in ws[header_row]:
        cell.fill = _fill(BRAND.BLUE)
        cell.font = Font(name=BRAND.fonts["heading"], bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name=BRAND.fonts["body"], color=BRAND.INK.lstrip("#"))
            if isinstance(cell.value, (int, float)):
                cell.number_format = number_format


def finish(wb, path):
    wb.save(path)
    load_workbook(path)
    return path
