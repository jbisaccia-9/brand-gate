"""Word builders. new_document() rewrites the built-in styles so plain
python-docx calls come out branded - no per-call font fiddling."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from brand import BRAND


def _rgb(hexstr):
    return RGBColor.from_string(hexstr.lstrip("#"))


def _style(doc, name, font, size, color, bold=False):
    st = doc.styles[name]
    st.font.name = font
    # East-Asian font slot must match or Word falls back silently.
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.font.size = Pt(size)
    st.font.color.rgb = _rgb(color)
    st.font.bold = bold


def new_document(title, subtitle="", date="", header="letterhead"):
    doc = Document()
    _style(doc, "Normal", BRAND.fonts["body"], 11, BRAND.INK)
    for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
        _style(doc, f"Heading {lvl}", BRAND.fonts["heading"], size, BRAND.BLUE)
    _style(doc, "Title", BRAND.fonts["heading"], 24, BRAND.BLUE)
    # Letterhead: logo + rule + title on page 1, no page break (memos/one-pagers).
    # Title page: same block followed by a page break (reports/proposals).
    doc.add_picture(BRAND.logo("primary"), width=Inches(1.6))
    _rule(doc.add_paragraph(), BRAND.AMBER)
    doc.add_paragraph(title, style="Title")
    if subtitle:
        doc.add_paragraph(subtitle)
    if date:
        doc.add_paragraph(date)
    if header == "titlepage":
        doc.add_page_break()
    return doc


def _rule(paragraph, color):
    """A bottom border on an empty paragraph = a horizontal rule Word keeps."""
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "12"), ("w:space", "1"), ("w:color", color.lstrip("#"))):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)


def _shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color.lstrip("#"))
    tcPr.append(shd)


def style_table(table):
    """Header row on a primary fill with white type; body rows on paper."""
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            _shade(cell, BRAND.BLUE if i == 0 else BRAND.PAPER)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = _rgb("#FFFFFF" if i == 0 else BRAND.INK)
                    r.font.bold = i == 0


def add_callout(doc, text):
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    _shade(c, BRAND.AMBER_TINT)
    c.paragraphs[0].add_run(text)          # inherits Normal -> INK, never tint-on-tint
    return t


def finish(doc, path):
    doc.save(path)
    Document(path)                          # reopen: malformed parts fail here, not on the recipient's machine
    return path
