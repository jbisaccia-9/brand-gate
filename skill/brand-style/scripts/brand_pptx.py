"""PowerPoint builders. Logo-bearing surfaces stay light; saturated fills
carry type only. chart_slide() makes a native, editable chart."""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from brand import BRAND

W, H = Inches(13.333), Inches(7.5)


def _rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


def _bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _text(slide, x, y, w, h, text, size, color, font=None, bold=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = _rgb(color)
    r.font.name = font or BRAND.fonts["body"]
    r.font.bold = bold
    return tb


def new_deck():
    d = Presentation()
    d.slide_width, d.slide_height = W, H
    return d


def _blank(d):
    return d.slides.add_slide(d.slide_layouts[6])


def title_slide(d, title, subtitle="", date=""):
    s = _blank(d)
    _bg(s, BRAND.PAPER)                                   # light: the logo lives here
    s.shapes.add_picture(BRAND.logo("primary"), Inches(0.6), Inches(0.5), width=Inches(2.2))
    band = s.shapes.add_shape(1, Inches(0), Inches(6.9), W, Inches(0.6))
    band.fill.solid(); band.fill.fore_color.rgb = _rgb(BRAND.GREEN); band.line.fill.background()
    _text(s, Inches(0.6), Inches(2.6), Inches(12), Inches(1.2), title, 40, BRAND.BLUE, BRAND.fonts["heading"])
    _text(s, Inches(0.6), Inches(3.8), Inches(12), Inches(0.8), subtitle, 20, BRAND.INK)
    _text(s, Inches(0.6), Inches(4.5), Inches(12), Inches(0.6), date, 14, BRAND.INK)
    return s


def section_slide(d, title, kicker=""):
    s = _blank(d)
    _bg(s, BRAND.GREEN)                                   # saturated: NO logo on this surface
    _text(s, Inches(0.8), Inches(2.8), Inches(11), Inches(1.2), title, 36, "#FFFFFF", BRAND.fonts["heading"])
    _text(s, Inches(0.8), Inches(2.2), Inches(11), Inches(0.6), kicker, 16, "#FFFFFF")
    return s


def content_slide(d, title):
    s = _blank(d)
    _bg(s, "#FFFFFF")
    _text(s, Inches(0.6), Inches(0.4), Inches(11), Inches(0.9), title, 28, BRAND.BLUE, BRAND.fonts["heading"])
    rule = s.shapes.add_shape(1, Inches(0.6), Inches(1.25), Inches(12.1), Inches(0.05))
    rule.fill.solid(); rule.fill.fore_color.rgb = _rgb(BRAND.AMBER); rule.line.fill.background()
    return s, Inches(1.5)


def bullets_slide(d, title, items):
    s, top = content_slide(d, title)
    tb = s.shapes.add_textbox(Inches(0.6), top, Inches(12), Inches(5))
    tf = tb.text_frame
    for i, item in enumerate(items):
        text, level = (item, 0) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        r = p.add_run(); r.text = ("• " if level == 0 else "– ") + text
        r.font.size = Pt(20 - 3 * level); r.font.color.rgb = _rgb(BRAND.INK); r.font.name = BRAND.fonts["body"]
    return s


def kpi_slide(d, title, kpis):
    s, top = content_slide(d, title)
    w = Inches(12) / max(len(kpis), 1)
    for i, (value, label) in enumerate(kpis):
        x = Inches(0.6) + w * i
        card = s.shapes.add_shape(1, x, top, w - Inches(0.3), Inches(2.4))
        card.fill.solid(); card.fill.fore_color.rgb = _rgb(BRAND.BLUE_TINT); card.line.fill.background()
        _text(s, x, top + Inches(0.3), w - Inches(0.3), Inches(1.2), value, 40, BRAND.BLUE, BRAND.fonts["heading"], True)
        _text(s, x, top + Inches(1.5), w - Inches(0.3), Inches(0.6), label, 14, BRAND.INK)
    return s


def chart_slide(d, title, categories, series):
    s, top = content_slide(d, title)
    data = CategoryChartData()
    data.categories = categories
    for name, values in series.items():
        data.add_series(name, values)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), top, Inches(12), Inches(5.2), data)
    for i, ser in enumerate(gf.chart.series):          # series order is the rule; never let Office pick
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = _rgb(BRAND.CHART_SERIES[i % len(BRAND.CHART_SERIES)])
    return s


def closing_slide(d, line="Thank you"):
    s = title_slide(d, line)
    return s


def finish(d, path):
    d.save(path)
    Presentation(path)
    return path
